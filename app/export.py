# -*- coding: utf-8 -*-
"""
匯出：標案清單 → Excel；收藏的截止日 → 行事曆（.ics，可匯入 Google 日曆／Outlook）。
"""

import hashlib
import io
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def tenders_xlsx(tenders: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "標案"
    headers = ["案名", "機關", "產業", "預算（元）", "截止投標", "履約地點", "招標方式", "標的分類", "公告連結"]
    ws.append(headers)
    for t in tenders:
        ws.append([t["title"], t["unit"], t["industries"], t["budget"], t["deadline"], t["location"],
                   t["tender_type"], t["category"], t["url"]])
        ws.cell(ws.max_row, 9).hyperlink = t["url"]
        ws.cell(ws.max_row, 9).style = "Hyperlink"

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4A90D9")
        cell.alignment = Alignment(horizontal="center")
    for col, width in zip(range(1, 10), [50, 28, 14, 14, 18, 18, 22, 28, 40]):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ics_escape(text: str) -> str:
    return (text or "").replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def _fold(line: str) -> list[str]:
    """iCalendar 規定每行最多 75 bytes，超過要折行（續行以空白開頭），且不能切斷 UTF-8 字元"""
    out, cur, size = [], "", 0
    for ch in line:
        n = len(ch.encode("utf-8"))
        if size + n > (75 if not out else 74):
            out.append(cur)
            cur, size = "", 0
        cur += ch
        size += n
    out.append(cur)
    return [out[0]] + [" " + x for x in out[1:]]


def deadlines_ics(tenders: list[dict]) -> bytes:
    """每個有截止日的標案建立一個事件（截止前 1 小時開始，前一天提醒）"""
    stamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    lines = [
        "BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//tender-subsidy-helper//ZH-TW", "CALSCALE:GREGORIAN",
        "X-WR-CALNAME:標案截止日", "X-WR-TIMEZONE:Asia/Taipei",
        "BEGIN:VTIMEZONE", "TZID:Asia/Taipei", "BEGIN:STANDARD", "DTSTART:19700101T000000",
        "TZOFFSETFROM:+0800", "TZOFFSETTO:+0800", "TZNAME:CST", "END:STANDARD", "END:VTIMEZONE",
    ]
    for t in tenders:
        if not t.get("deadline"):
            continue
        end = datetime.strptime(t["deadline"], "%Y-%m-%d %H:%M")
        start = end - timedelta(hours=1)
        # UID 要固定，重複匯入時行事曆才會更新同一個事件而不是重複新增
        uid = f"{hashlib.md5(t['url'].encode()).hexdigest()}@tender-subsidy-helper"
        desc = f"機關：{t['unit']}\n預算：{t.get('budget') or '未公開'}\n公告：{t['url']}"
        lines += [
            "BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{stamp}",
            f"DTSTART;TZID=Asia/Taipei:{start:%Y%m%dT%H%M%S}",
            f"DTEND;TZID=Asia/Taipei:{end:%Y%m%dT%H%M%S}",
            f"SUMMARY:{_ics_escape('⏰ 截止投標：' + t['title'])}",
            f"DESCRIPTION:{_ics_escape(desc)}",
            f"URL:{t['url']}",
            "BEGIN:VALARM", "ACTION:DISPLAY", "TRIGGER:-P1D",
            f"DESCRIPTION:{_ics_escape('明天截止：' + t['title'])}", "END:VALARM",
            "END:VEVENT",
        ]
    lines.append("END:VCALENDAR")
    folded = [part for line in lines for part in _fold(line)]
    return ("\r\n".join(folded) + "\r\n").encode("utf-8")
